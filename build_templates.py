"""Build the Neudata Word templates.

Produces .docx sources in ./build; build.ps1 then uses Microsoft Word to save
them as .dotx templates and render PDF previews.

    python build_templates.py

Placeholders are written as [Bracketed text] in grey italics so they are easy
to find (Ctrl+H for "[") and impossible to miss before sending.
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
LOGO = HERE / "assets" / "neudata-logo.png"
OUT = HERE / "build"
OUT.mkdir(exist_ok=True)

# ----- Brand -------------------------------------------------------------------
TEAL, NAVY, BLUE = "055F56", "04242F", "0B376C"
MIST, GREY, INK, SEAFOAM = "EAF2F5", "8A99A3", "1F2D33", "9FD3CB"
HEAD_FONT = "Century Gothic"   # as in the Neudata PowerPoint template and previous invoice
BODY_FONT = "Arial"            # safe on Windows, Mac and Word Online

COMPANY = "Neudata Consulting Ltd"
TAGLINE = "Insight. Impact. Innovation."
WEB, EMAIL = "www.neu-data.com", "contact@neu-data.com"
OFFICES = "Nairobi, Kenya  ·  Antwerp, Belgium"


def rgb(hexval):
    return RGBColor.from_string(hexval)


# ----- Low-level XML helpers ------------------------------------------------------
def set_fonts(rpr_owner, name):
    rpr = rpr_owner.get_or_add_rPr() if hasattr(rpr_owner, "get_or_add_rPr") else rpr_owner
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def style_font(style, name=BODY_FONT, size=None, color=None, bold=None, italic=None):
    style.font.name = name
    set_fonts(style.element, name)
    if size is not None:
        style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


# Elements that must come AFTER a given child, per the WordprocessingML schema.
# Used with insert_element_before() so every element lands in schema order.
PPR_AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                  "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                  "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                  "w:pPrChange")
PPR_AFTER_SHD = PPR_AFTER_PBDR[1:]
TCPR_AFTER_BORDERS = ("w:shd", "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign",
                      "w:hideMark", "w:headers", "w:cellIns", "w:cellDel", "w:cellMerge", "w:tcPrChange")
TCPR_AFTER_SHD = TCPR_AFTER_BORDERS[1:]
TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                       "w:tblDescription", "w:tblPrChange")
TBLPR_AFTER_CELLMAR = ("w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")
TRPR_AFTER_CANTSPLIT = ("w:trHeight", "w:tblHeader", "w:tblCellSpacing", "w:jc", "w:hidden",
                        "w:ins", "w:del", "w:trPrChange")
TRPR_AFTER_TBLHEADER = ("w:tblCellSpacing", "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange")


def para_border(paragraph_or_style, side="bottom", color=TEAL, size=8, space=4):
    el = paragraph_or_style.element if hasattr(paragraph_or_style, "element") else paragraph_or_style._p
    ppr = el.get_or_add_pPr()
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        ppr.insert_element_before(bdr, *PPR_AFTER_PBDR)
    b = OxmlElement(f"w:{side}")
    for k, v in dict(val="single", sz=str(size), space=str(space), color=color).items():
        b.set(qn(f"w:{k}"), v)
    # border sides are ordered top, left, bottom, right
    order = ["top", "left", "bottom", "right", "between", "bar"]
    later = {qn(f"w:{s}") for s in order[order.index(side) + 1:]}
    successor = next((child for child in bdr if child.tag in later), None)
    if successor is not None:
        successor.addprevious(b)
    else:
        bdr.append(b)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), fill)
    tcpr.insert_element_before(s, *TCPR_AFTER_SHD)


def cell_borders(cell, color="D5DDE1", size=4, sides=("top", "bottom", "left", "right")):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        if side in sides:
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), str(size))
            b.set(qn("w:color"), color)
        else:
            b.set(qn("w:val"), "nil")
        borders.append(b)
    tcpr.insert_element_before(borders, *TCPR_AFTER_BORDERS)


def cell_margins(table, top=80, bottom=80, left=110, right=110):
    tblpr = table._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tblpr.insert_element_before(m, *TBLPR_AFTER_CELLMAR)


def no_table_borders(table):
    tblpr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.insert_element_before(b, *TBLPR_AFTER_BORDERS)


def col_widths(table, widths_cm):
    table.autofit = False
    tblgrid = table._tbl.tblGrid
    for gc, w in zip(tblgrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths_cm):
            cell.width = Cm(w)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trpr.insert_element_before(h, *TRPR_AFTER_TBLHEADER)


def add_field(paragraph, code, size=8, color=GREY):
    """Insert a Word field (PAGE, NUMPAGES, DATE...)."""
    run = paragraph.add_run()
    for kind, text in (("begin", None), ("instr", code), ("separate", None), ("text", "1"), ("end", None)):
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {code} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    return run


def placeholder(paragraph, text):
    """Grey italic [placeholder] text."""
    r = paragraph.add_run(f"[{text}]")
    r.italic = True
    r.font.color.rgb = rgb("6B7B84")
    return r


def text(paragraph, s, bold=None, color=None, size=None, italic=None, font=None):
    r = paragraph.add_run(s)
    if bold is not None:
        r.bold = bold
    if italic is not None:
        r.italic = italic
    if color:
        r.font.color.rgb = rgb(color)
    if size:
        r.font.size = Pt(size)
    if font:
        r.font.name = font
        set_fonts(r._r, font)
    return r


# ----- Base document: styles, page, header, footer ------------------------------------
def base_document(doc_label, first_page_letterhead=False):
    doc = Document()
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2.4)
    sec.header_distance, sec.footer_distance = Cm(1.0), Cm(1.0)

    st = doc.styles
    normal = st["Normal"]
    style_font(normal, BODY_FONT, 10.5, INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    style_font(st["Title"], HEAD_FONT, 26, NAVY, True)
    st["Title"].paragraph_format.space_after = Pt(6)
    title_ppr = st["Title"].element.get_or_add_pPr()
    for b in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(b)
    style_font(st["Subtitle"], HEAD_FONT, 14, TEAL, False, False)

    for name, size, color, before in (("Heading 1", 15, NAVY, 18), ("Heading 2", 12.5, BLUE, 12),
                                      ("Heading 3", 11, TEAL, 10)):
        s = st[name]
        style_font(s, HEAD_FONT, size, color, True, False)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    para_border(st["Heading 1"], "bottom", TEAL, 10, 4)

    for name in ("List Bullet", "List Number"):
        style_font(st[name], BODY_FONT, 10.5, INK)
    style_font(st["Caption"], BODY_FONT, 9, NAVY, False, True)

    # Neudata paragraph styles
    kf = st.add_style("Key Findings", WD_STYLE_TYPE.PARAGRAPH)
    kf.base_style = normal
    style_font(kf, BODY_FONT, 10.5, NAVY)
    ppr = kf.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), MIST)
    ppr.insert_element_before(shd, *PPR_AFTER_SHD)
    para_border(kf, "left", TEAL, 30, 8)
    kf.paragraph_format.left_indent = Cm(0.35)

    small = st.add_style("Neudata Small", WD_STYLE_TYPE.PARAGRAPH)
    small.base_style = normal
    style_font(small, BODY_FONT, 8.5, GREY)
    small.paragraph_format.space_after = Pt(0)

    label = st.add_style("Neudata Label", WD_STYLE_TYPE.CHARACTER)
    style_font(label, HEAD_FONT, None, NAVY, True)

    # Word's Header/Footer styles carry centre and right tab stops for Letter paper;
    # replace them with a single right tab at the A4 text width.
    for name in ("Header", "Footer"):
        if name in [s.name for s in st]:
            ts = st[name].paragraph_format.tab_stops
            ts.clear_all()
            ts.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
            style_font(st[name], BODY_FONT)

    # Header: document label left, logo right, teal rule. Footer: tagline + page n / N.
    header = sec.header
    hp = header.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    text(hp, doc_label, color=NAVY, size=8)
    hp.add_run("\t").add_picture(str(LOGO), height=Cm(0.95))
    para_border(hp, "bottom", TEAL, 6, 4)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT)
    text(fp, f"{COMPANY} · ", color=GREY, size=8)
    text(fp, TAGLINE, color=GREY, size=8, italic=True)
    fp.add_run("\t")
    add_field(fp, "PAGE")
    text(fp, " / ", color=GREY, size=8)
    add_field(fp, "NUMPAGES")

    if first_page_letterhead:
        sec.different_first_page_header_footer = True
        letterhead_first_page(sec)
    return doc


def letterhead_first_page(sec):
    """Large logo + company block on page one; company details in the footer."""
    h = sec.first_page_header
    tbl = h.add_table(rows=1, cols=2, width=Cm(16.6))
    no_table_borders(tbl)
    col_widths(tbl, [7.0, 9.6])
    left, right = tbl.rows[0].cells
    lp = left.paragraphs[0]
    lp.add_run().add_picture(str(LOGO), height=Cm(2.0))
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text(rp, COMPANY, bold=True, color=NAVY, size=13, font=HEAD_FONT)
    for line, colr in (("Data Science · Biostatistics · Research Analytics", TEAL),
                       (OFFICES, INK), (f"{EMAIL}  ·  {WEB}", BLUE)):
        p = right.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(0)
        text(p, line, color=colr, size=8.5)
    # remove the default empty paragraph before the table, add a teal rule after it
    first = h.paragraphs[0]
    para_border(first, "bottom", TEAL, 10, 1)
    h._element.remove(first._p)
    h._element.append(first._p)

    f = sec.first_page_footer
    fp = f.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_border(fp, "top", TEAL, 6, 6)
    text(fp, f"{COMPANY}  ·  ", color=NAVY, size=8, bold=True)
    placeholder(fp, "Company registration no.")
    text(fp, "  ·  PIN/VAT ", color=GREY, size=8)
    placeholder(fp, "number")
    for r in fp.runs:
        r.font.size = Pt(8)
    p2 = f.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p2, f"{OFFICES}  ·  {EMAIL}  ·  {WEB}", color=GREY, size=8)
    p3 = f.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p3, TAGLINE, color=TEAL, size=8, italic=True)


def brand_table(doc, header, rows, widths, align_right=(), total_rows=0):
    """Navy header row, light borders, striped body, optional bold total rows."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_margins(t)
    col_widths(t, widths)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        shade(c, NAVY)
        cell_borders(c, NAVY)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if i in align_right:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        text(p, h, bold=True, color="FFFFFF", size=9.5)
    repeat_header(t.rows[0])
    n_body = len(rows) - total_rows
    # Keep the table on one page: rows never split, and each row stays with the next
    for r_i, tr in enumerate(t.rows):
        trpr = tr._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trpr.insert_element_before(cs, *TRPR_AFTER_CANTSPLIT)
        if r_i < len(t.rows) - 1:
            for c in tr.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True
    for r_i, row in enumerate(rows, start=1):
        is_total = r_i > n_body
        for c_i, val in enumerate(row):
            c = t.rows[r_i].cells[c_i]
            cell_borders(c, "D5DDE1")
            if is_total:
                shade(c, MIST)
            elif r_i % 2 == 0:
                shade(c, "F5F9FA")
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if c_i in align_right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if isinstance(val, tuple) and val[0] == "ph":
                placeholder(p, val[1])
            else:
                text(p, str(val), bold=is_total or None, size=9.5)
    return t


def spacer(doc, pts=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(pts)
    return p


def kv_block(doc, items, label_width=4.2):
    """Two-column label/value block without borders (e.g. invoice details)."""
    t = doc.add_table(rows=len(items), cols=2)
    no_table_borders(t)
    cell_margins(t, 30, 30, 0, 80)
    col_widths(t, [label_width, 16.6 - label_width])
    for (label, value), row in zip(items, t.rows):
        lp = row.cells[0].paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        text(lp, label, bold=True, color=NAVY, size=10)
        vp = row.cells[1].paragraphs[0]
        vp.paragraph_format.space_after = Pt(0)
        if isinstance(value, tuple):
            placeholder(vp, value[1])
        else:
            text(vp, value, size=10)
    return t


def guidance(doc, s):
    p = doc.add_paragraph()
    r = p.add_run(s)
    r.italic = True
    r.font.color.rgb = rgb("6B7B84")
    r.font.size = Pt(9.5)
    return p


# =====================================================================================
# 1. General document template
# =====================================================================================
def build_document():
    doc = base_document("Document title")
    doc.add_paragraph("Document title", style="Title")
    doc.add_paragraph("Subtitle or short description", style="Subtitle")
    p = doc.add_paragraph(style="Neudata Small")
    placeholder(p, "Author")
    text(p, "  ·  ")
    placeholder(p, "DD-MM-YYYY")
    text(p, "  ·  Version ")
    placeholder(p, "1.0")
    spacer(doc, 14)

    doc.add_heading("1  Heading 1", level=1)
    doc.add_paragraph(
        "Body text uses the Normal style (Arial 10.5 pt). Headings use Century Gothic in the "
        "Neudata colours: navy with a teal rule for Heading 1, logo blue for Heading 2 and teal "
        "for Heading 3. Use the Styles gallery (Home → Styles) rather than manual formatting.")
    doc.add_paragraph("Key Findings style: highlight the main message of a section in a tinted box.",
                      style="Key Findings")
    doc.add_heading("1.1  Heading 2", level=2)
    for item in ("Bulleted list item (List Bullet style)", "Second item", "Third item"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Heading 3", level=3)
    doc.add_paragraph("Tables use a navy header row and light striping:")
    brand_table(doc, ["Item", "Description", "Value"],
                [["A", "First row", "1,200"], ["B", "Second row", "640"], ["C", "Third row", "560"]],
                [3.0, 9.6, 4.0], align_right=(2,))
    cap = doc.add_paragraph("Table 1: Example table caption (Caption style).", style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.core_properties.title = "Neudata document"
    doc.core_properties.author = COMPANY
    doc.save(OUT / "Neudata-Document.docx")


# =====================================================================================
# 2. Letterhead
# =====================================================================================
def build_letterhead():
    doc = base_document("Letter", first_page_letterhead=True)
    spacer(doc, 10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    text(p, "Date: ", bold=True, color=NAVY)
    placeholder(p, "DD Month YYYY")
    p = doc.add_paragraph()
    text(p, "Our ref: ", bold=True, color=NAVY)
    placeholder(p, "NDC-YYYY-000")
    spacer(doc, 8)
    for line in ("Recipient name", "Title, Organisation", "Address line", "City, Country"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        placeholder(p, line)
    spacer(doc, 12)
    p = doc.add_paragraph()
    text(p, "Dear ")
    placeholder(p, "Name")
    text(p, ",")
    p = doc.add_paragraph()
    text(p, "RE: ", bold=True, color=NAVY)
    r = placeholder(p, "SUBJECT OF THE LETTER")
    r.bold = True
    for _ in range(2):
        p = doc.add_paragraph()
        placeholder(p, "Body of the letter.")
    doc.add_paragraph("Yours sincerely,")
    spacer(doc, 30)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    text(p, "_______________________________", color=GREY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = placeholder(p, "Name")
    r.bold = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    placeholder(p, "Title")
    p = doc.add_paragraph()
    text(p, COMPANY, color=TEAL, bold=True)
    doc.core_properties.title = "Neudata letterhead"
    doc.core_properties.author = COMPANY
    doc.save(OUT / "Neudata-Letterhead.docx")


# =====================================================================================
# 3. Invoice
# =====================================================================================
def build_invoice():
    doc = base_document("Invoice", first_page_letterhead=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    text(p, "INVOICE", bold=True, color=NAVY, size=22, font=HEAD_FONT)

    # Invoice details (left) and Bill to / Bill from (two columns)
    kv_block(doc, [
        ("Invoice number:", ("ph", "NDCL000/CC  e.g. NDCL002/UG")),
        ("Invoice date:", ("ph", "DD Month YYYY")),
        ("Due date:", ("ph", "DD Month YYYY — e.g. 30 days from invoice date")),
        ("Contract / PO ref:", ("ph", "Contract or purchase order reference")),
    ])
    spacer(doc, 6)

    t = doc.add_table(rows=1, cols=2)
    no_table_borders(t)
    cell_margins(t, 60, 60, 0, 120)
    col_widths(t, [8.3, 8.3])
    for cell, heading, lines in (
        (t.rows[0].cells[0], "Bill to", ["Client organisation", "Attention: contact name", "Street address",
                                         "P. O. Box", "City, Country"]),
        (t.rows[0].cells[1], "Bill from", None),
    ):
        hp = cell.paragraphs[0]
        hp.paragraph_format.space_after = Pt(2)
        text(hp, heading, bold=True, color=TEAL, size=10.5, font=HEAD_FONT)
        para_border(hp, "bottom", TEAL, 6, 1)
        if lines:
            for line in lines:
                q = cell.add_paragraph()
                q.paragraph_format.space_after = Pt(0)
                placeholder(q, line)
        else:
            q = cell.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            text(q, COMPANY, bold=True)
            for line in ("Street address / P. O. Box", "City, Country", "PIN / VAT number"):
                q = cell.add_paragraph()
                q.paragraph_format.space_after = Pt(0)
                placeholder(q, line)
            q = cell.add_paragraph()
            q.paragraph_format.space_after = Pt(0)
            text(q, EMAIL, color=BLUE)
    spacer(doc, 12)

    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    text(p, "Payment description: ", bold=True, color=NAVY)
    placeholder(p, "Short description of the assignment, e.g. analysis and writing of four papers on ...")

    brand_table(
        doc,
        ["Description", "Quantity", "Unit price", "Amount (USD)"],
        [
            [("ph", "Consultancy payment: deliverable or milestone, e.g. 50% as agreed in the contract"),
             ("ph", "e.g. 2 of 4"), ("ph", "0.00"), ("ph", "0.00")],
            [("ph", "Additional line item"), ("ph", "1"), ("ph", "0.00"), ("ph", "0.00")],
            ["Subtotal", "", "", ("ph", "0.00")],
            ["Tax / VAT (if applicable)", "", "", ("ph", "0.00")],
            ["TOTAL DUE (USD)", "", "", ("ph", "0.00")],
        ],
        [8.6, 2.6, 2.6, 2.8], align_right=(1, 2, 3), total_rows=3,
    )
    spacer(doc, 8)
    p = doc.add_paragraph()
    text(p, "Amount in words: ", bold=True, color=NAVY)
    r = placeholder(p, "One thousand, one hundred and sixty-six US dollars ONLY")
    r.underline = True
    text(p, " or equivalent in ")
    placeholder(p, "local currency")
    text(p, ".")

    # Payment details box (left) and signature (right), side by side
    spacer(doc, 4)
    outer = doc.add_table(rows=1, cols=2)
    no_table_borders(outer)
    cell_margins(outer, 0, 0, 0, 0)
    col_widths(outer, [9.8, 6.8])
    box_cell, sig_cell = outer.rows[0].cells
    box = box_cell.add_table(rows=1, cols=1)
    box_cell._tc.remove(box_cell.paragraphs[0]._p)
    box_cell._tc.append(OxmlElement("w:p"))   # a cell must end with a paragraph
    cell_margins(box, 90, 90, 160, 160)
    col_widths(box, [9.6])
    c = box.rows[0].cells[0]
    shade(c, MIST)
    cell_borders(c, MIST)
    hp = c.paragraphs[0]
    hp.paragraph_format.space_after = Pt(4)
    text(hp, "Payment to be made to", bold=True, color=TEAL, font=HEAD_FONT)
    inner = c.add_table(rows=7, cols=2)
    no_table_borders(inner)
    cell_margins(inner, 10, 10, 0, 60)
    col_widths(inner, [3.3, 6.0])
    for row, (label, ph) in zip(inner.rows, [
        ("Bank name:", "Bank"), ("Branch:", "Branch"), ("SWIFT code:", "SWIFT/BIC"),
        ("Bank code:", "Bank code"), ("Branch code:", "Branch code"),
        ("Account name:", "Account name"), ("Account number:", "Account number"),
    ]):
        lp = row.cells[0].paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        text(lp, label, color=NAVY, size=9.5)
        vp = row.cells[1].paragraphs[0]
        vp.paragraph_format.space_after = Pt(0)
        r = placeholder(vp, ph)
        r.font.size = Pt(9.5)
    ip = c.add_paragraph()
    ip.paragraph_format.space_before = Pt(4)
    ip.paragraph_format.space_after = Pt(0)
    text(ip, "Please quote the invoice number as the payment reference.", color=GREY, size=8.5, italic=True)

    sig_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    sp = sig_cell.paragraphs[0]
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.space_after = Pt(0)
    text(sp, "_____________________________", color=GREY)
    for line, kind in (("Name, Title", "ph"), (f"For and on behalf of {COMPANY}", "txt")):
        q = sig_cell.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(0)
        if kind == "ph":
            placeholder(q, line)
        else:
            text(q, line, color=GREY, size=8.5)
    spacer(doc, 4)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    text(p, "Thank you for your business.", bold=True, color=TEAL, font=HEAD_FONT, size=11)
    doc.core_properties.title = "Neudata invoice"
    doc.core_properties.author = COMPANY
    doc.save(OUT / "Neudata-Invoice.docx")


# =====================================================================================
# 4. Proposal
# =====================================================================================
def cover(doc, kind, title_ph, subtitle_ph):
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True  # clean cover page
    p = doc.add_paragraph()
    p.add_run().add_picture(str(LOGO), height=Cm(3.0))
    spacer(doc, 40)
    p = doc.add_paragraph()
    text(p, kind.upper(), bold=True, color=TEAL, size=10, font=HEAD_FONT)
    band = doc.add_table(rows=1, cols=1)
    cell_margins(band, 420, 420, 360, 360)
    col_widths(band, [16.6])
    c = band.rows[0].cells[0]
    shade(c, NAVY)
    cell_borders(c, NAVY)
    tp = c.paragraphs[0]
    tp.paragraph_format.space_after = Pt(8)
    r = text(tp, f"[{title_ph}]", bold=True, color="FFFFFF", size=24, font=HEAD_FONT)
    sp = c.add_paragraph()
    sp.paragraph_format.space_after = Pt(14)
    text(sp, f"[{subtitle_ph}]", color=SEAFOAM, size=13)
    cp = c.add_paragraph()
    cp.paragraph_format.space_after = Pt(0)
    text(cp, "Prepared for: ", color="FFFFFF", size=10.5)
    text(cp, "[Client organisation]", color="FFFFFF", size=10.5, bold=True)
    spacer(doc, 120)
    kv_block(doc, [
        ("Submitted by:", COMPANY),
        ("Contact:", ("ph", "Name, Title · email · phone")),
        ("Date:", ("ph", "DD Month YYYY")),
        ("Reference:", ("ph", "NDC-PROP-YYYY-000")),
        ("Validity:", ("ph", "This proposal is valid for 60 days")),
    ], label_width=3.6)
    spacer(doc, 10)
    p = doc.add_paragraph(style="Neudata Small")
    text(p, f"{COMPANY}  ·  {OFFICES}  ·  {EMAIL}  ·  {WEB}")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_proposal():
    doc = base_document("Technical and financial proposal")
    cover(doc, "Technical and financial proposal", "Proposal title", "Assignment or study name")

    # Looks like Heading 1 but is not one, so it does not list itself in the contents
    ch = doc.add_paragraph()
    ch.paragraph_format.space_after = Pt(10)
    text(ch, "Contents", bold=True, color=NAVY, size=15, font=HEAD_FONT)
    para_border(ch, "bottom", TEAL, 10, 4)
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, val in (("begin", None), ("instr", 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                      ("text", "Right-click and choose Update Field to build the table of contents."),
                      ("end", None)):
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); run._r.append(fc)
        elif kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = val; run._r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = val; run._r.append(t)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading("1  Executive summary", level=1)
    guidance(doc, "Half a page: the client's need, what Neudata proposes, the main deliverables, "
                  "timeline and total cost. Write this last.")
    doc.add_paragraph("Key message: what the client gets, by when, and why Neudata.", style="Key Findings")

    doc.add_heading("2  Understanding of the assignment", level=1)
    guidance(doc, "Restate the background and the problem in your own words, showing you have read "
                  "the terms of reference. Note any assumptions or clarifications.")

    doc.add_heading("3  Objectives and scope", level=1)
    doc.add_heading("3.1  Objectives", level=2)
    for i in range(1, 4):
        p = doc.add_paragraph(style="List Number")
        placeholder(p, f"Objective {i}")
    doc.add_heading("3.2  Scope", level=2)
    guidance(doc, "What is included and — just as important — what is out of scope.")

    doc.add_heading("4  Approach and methodology", level=1)
    guidance(doc, "Study design, data sources, sampling and sample size, data management and quality "
                  "assurance, statistical analysis, de-identification and data protection, reproducibility.")
    for sub in ("4.1  Data management and quality assurance", "4.2  Statistical analysis",
                "4.3  Data protection and ethics", "4.4  Capacity building and knowledge transfer"):
        doc.add_heading(sub, level=2)
        p = doc.add_paragraph()
        placeholder(p, "Describe")

    doc.add_heading("5  Deliverables", level=1)
    brand_table(doc, ["#", "Deliverable", "Format", "Due"],
                [["1", ("ph", "Inception report and statistical analysis plan"), ("ph", "Word / PDF"),
                  ("ph", "Week 2")],
                 ["2", ("ph", "Cleaned, de-identified analysis dataset"), ("ph", "CSV / Stata / R"),
                  ("ph", "Week 6")],
                 ["3", ("ph", "Final report and presentation"), ("ph", "Word / PDF / PowerPoint"),
                  ("ph", "Week 10")]],
                [1.0, 9.0, 3.6, 3.0])

    doc.add_heading("6  Work plan and timeline", level=1)
    brand_table(doc, ["Phase", "Activities", "Weeks"],
                [[("ph", "Inception"), ("ph", "Kick-off, document review, analysis plan"), ("ph", "1–2")],
                 [("ph", "Data"), ("ph", "Data cleaning, quality checks, de-identification"), ("ph", "3–6")],
                 [("ph", "Analysis"), ("ph", "Statistical analysis, tables and figures"), ("ph", "6–8")],
                 [("ph", "Reporting"), ("ph", "Draft report, review, final report"), ("ph", "8–10")]],
                [3.2, 10.4, 3.0])

    doc.add_heading("7  About Neudata and the team", level=1)
    doc.add_paragraph(
        f"{COMPANY} is a data science, biostatistics and research analytics consultancy that helps "
        "organisations, researchers and decision-makers turn complex data into clear, credible and "
        "actionable evidence — from study design and data management to statistical analysis, "
        "reporting and capacity development.")
    brand_table(doc, ["Name", "Role on this assignment", "Relevant expertise"],
                [[("ph", "Name"), ("ph", "Lead biostatistician"), ("ph", "Expertise")],
                 [("ph", "Name"), ("ph", "Data manager"), ("ph", "Expertise")],
                 [("ph", "Name"), ("ph", "Analyst"), ("ph", "Expertise")]],
                [4.2, 5.4, 7.0])

    doc.add_heading("8  Budget", level=1)
    brand_table(doc, ["Item", "Unit", "Quantity", "Rate (USD)", "Amount (USD)"],
                [[("ph", "Lead biostatistician"), ("ph", "day"), ("ph", "0"), ("ph", "0.00"), ("ph", "0.00")],
                 [("ph", "Data manager"), ("ph", "day"), ("ph", "0"), ("ph", "0.00"), ("ph", "0.00")],
                 [("ph", "Software, travel or other costs"), ("ph", "lump sum"), ("ph", "1"), ("ph", "0.00"),
                  ("ph", "0.00")],
                 ["Subtotal", "", "", "", ("ph", "0.00")],
                 ["Tax / VAT (if applicable)", "", "", "", ("ph", "0.00")],
                 ["TOTAL (USD)", "", "", "", ("ph", "0.00")]],
                [6.4, 2.2, 2.2, 2.8, 3.0], align_right=(2, 3, 4), total_rows=3)

    doc.add_heading("9  Payment schedule and terms", level=1)
    brand_table(doc, ["Milestone", "Share", "Amount (USD)"],
                [[("ph", "On signature of contract / inception report"), ("ph", "50%"), ("ph", "0.00")],
                 [("ph", "On approval of the final report"), ("ph", "50%"), ("ph", "0.00")]],
                [10.6, 2.4, 3.6], align_right=(1, 2))
    spacer(doc, 6)
    for term in ("Invoices are payable within 30 days of receipt.",
                 "Data are handled confidentially and in line with applicable data-protection law; "
                 "raw identifiable data are not retained after the assignment unless agreed.",
                 "Intellectual property in the deliverables transfers to the client on full payment; "
                 "Neudata retains its pre-existing tools, code libraries and methods.",
                 "This proposal is valid for 60 days from the date above."):
        doc.add_paragraph(term, style="List Bullet")

    doc.add_heading("10  Acceptance", level=1)
    doc.add_paragraph("Signed for and on behalf of the parties:")
    sig = doc.add_table(rows=4, cols=2)
    no_table_borders(sig)
    cell_margins(sig, 60, 60, 0, 200)
    col_widths(sig, [8.3, 8.3])
    for col, who in enumerate((COMPANY, "[Client organisation]")):
        cells = [sig.rows[i].cells[col] for i in range(4)]
        text(cells[0].paragraphs[0], who, bold=True, color=NAVY)
        cells[1].paragraphs[0].paragraph_format.space_before = Pt(24)
        text(cells[1].paragraphs[0], "Signature: ______________________", color=GREY)
        text(cells[2].paragraphs[0], "Name: ___________________________", color=GREY)
        text(cells[3].paragraphs[0], "Date: ____________________________", color=GREY)

    doc.core_properties.title = "Neudata proposal"
    doc.core_properties.author = COMPANY
    doc.save(OUT / "Neudata-Proposal.docx")


if __name__ == "__main__":
    build_document()
    build_letterhead()
    build_invoice()
    build_proposal()
    print("built:", sorted(p.name for p in OUT.glob("*.docx")))
